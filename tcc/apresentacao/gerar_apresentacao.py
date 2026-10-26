"""Gera Apresentacao_TCC.pptx a partir do conteúdo do artigo (tcc/artigo/artigo.tex).

Estilo baseado no template oficial PUC Campinas (visto em
exemplos_slides/EffAttNet_Glaucoma_Slides.pdf, fornecido pelo professor):
fundo cinza-claro, títulos serifados azul-marinho, subtítulos em negrito
azul, corpo em bullets pretos com termos-chave em negrito.

Todo número e afirmação técnica abaixo vem do artigo atual; nada foi
inventado para a apresentação.
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

# ── Paleta e tipografia (amostradas do exemplo oficial) ──────────────────────
BG      = RGBColor(0xEE, 0xEE, 0xF5)
NAVY    = RGBColor(0x14, 0x20, 0x5E)   # títulos serifados
BLUE    = RGBColor(0x1A, 0x1F, 0x9C)   # subtítulos em negrito
BLACK   = RGBColor(0x18, 0x18, 0x18)
GRAY    = RGBColor(0x6B, 0x6B, 0x76)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
TABLE_HDR_BG = RGBColor(0xDD, 0xDF, 0xEC)
TABLE_ALT_BG = RGBColor(0xF6, 0xF7, 0xFB)
ACCENT_YELLOW = RGBColor(0xC8, 0xB4, 0x00)  # mesma família do módulo de ML no diagrama

SERIF = "Georgia"
SANS  = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN_L
TITLE_TOP = Inches(0.45)
SUB_TOP = Inches(1.28)
BODY_TOP = Inches(1.95)

ASSETS = "assets"
DATA_LABEL = "Outubro de 2026"

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ── Helpers ───────────────────────────────────────────────────────────────────
def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # manda pro fundo
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def _set_run(run, text, size, color, bold=False, italic=False, font=SANS):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def _add_runs_markup(paragraph, text, size, color, font=SANS, base_italic=False):
    """Divide `text` em runs, tratando **negrito** e _itálico_ simples."""
    tokens = re.split(r"(\*\*.*?\*\*|_.*?_)", text)
    for tok in tokens:
        if not tok:
            continue
        run = paragraph.add_run()
        if tok.startswith("**") and tok.endswith("**"):
            _set_run(run, tok[2:-2], size, color, bold=True, italic=base_italic, font=font)
        elif tok.startswith("_") and tok.endswith("_"):
            _set_run(run, tok[1:-1], size, color, bold=False, italic=True, font=font)
        else:
            _set_run(run, tok, size, color, bold=False, italic=base_italic, font=font)


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box, tf


def add_title(slide, text, size=32, top=TITLE_TOP, left=MARGIN_L, width=None):
    width = width or (CONTENT_W)
    box, tf = add_textbox(slide, left, top, width, Inches(1.0))
    p = tf.paragraphs[0]
    _add_runs_markup(p, text, size, NAVY, font=SERIF)
    for r in p.runs:
        r.font.bold = True
    return box


def add_subheading(slide, text, top=SUB_TOP, left=MARGIN_L, size=18, width=None):
    width = width or CONTENT_W
    box, tf = add_textbox(slide, left, top, width, Inches(0.5))
    p = tf.paragraphs[0]
    _add_runs_markup(p, text, size, BLUE, font=SANS)
    for r in p.runs:
        r.font.bold = True
    return box


def add_bullets(slide, items, left=MARGIN_L, top=BODY_TOP, width=CONTENT_W,
                 height=Inches(5.0), size=16, space_after=10, bullet_char="•",
                 line_spacing=1.08):
    # nunca deixa a caixa nominal passar do rodapé do slide
    height = min(height, SLIDE_H - top - Inches(0.3))
    box, tf = add_textbox(slide, left, top, width, height)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        indent = 0
        text = item
        if isinstance(item, tuple):
            indent, text = item
        prefix = ("   " * indent) + (bullet_char + "  " if bullet_char else "")
        if prefix:
            run0 = p.add_run()
            _set_run(run0, prefix, size, BLACK if indent == 0 else GRAY, font=SANS)
        _add_runs_markup(p, text, size - (2 if indent else 0),
                          BLACK if indent == 0 else GRAY, font=SANS)
    return box


def add_page_chrome(slide, page_num, show_date=True):
    if show_date:
        box, tf = add_textbox(slide, SLIDE_W - Inches(3.2), Inches(0.28), Inches(2.9), Inches(0.35))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        _add_runs_markup(p, DATA_LABEL, 11, GRAY, font=SANS)
    box, tf = add_textbox(slide, SLIDE_W - Inches(1.0), SLIDE_H - Inches(0.5), Inches(0.7), Inches(0.3))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _add_runs_markup(p, str(page_num), 11, GRAY, font=SANS)


def content_slide(title, page_num, subheading=None, sub_top=SUB_TOP, body_top=None):
    s = new_slide()
    add_title(s, title)
    if subheading:
        add_subheading(s, subheading, top=sub_top)
        bt = body_top or (sub_top + Inches(0.55))
    else:
        bt = body_top or Inches(1.65)
    add_page_chrome(s, page_num)
    return s, bt


def add_picture_fit(slide, path, left, top, max_w, max_h, align="center"):
    # nunca deixa a imagem passar do rodapé do slide
    bottom_margin = Inches(0.3)
    available_h = SLIDE_H - top - bottom_margin
    max_h = min(max_h, available_h)
    im = Image.open(path)
    iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    if align == "center":
        left = left + (max_w - w) // 2
    top = top + (max_h - h) // 2
    slide.shapes.add_picture(path, left, top, width=w, height=h)
    return w, h


def style_table(table, n_cols, header_size=13, body_size=12.5, header_align=PP_ALIGN.CENTER):
    for ci in range(n_cols):
        table.columns[ci].width = table.columns[ci].width  # no-op, widths set by caller
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HDR_BG
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_ALT_BG if ri % 2 == 0 else WHITE
            tf = cell.text_frame
            for p in tf.paragraphs:
                p.alignment = header_align if ri == 0 else PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.size = Pt(header_size if ri == 0 else body_size)
                    r.font.bold = ri == 0
                    r.font.name = SANS
                    r.font.color.rgb = NAVY if ri == 0 else BLACK


def set_cell_text(cell, text, size=12.5, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _add_runs_markup(p, text, size, color, font=SANS)
    for r in p.runs:
        r.font.bold = bold


# ============================================================================
# SLIDE 1 — Capa
# ============================================================================
s = new_slide()
logo = s.shapes.add_picture(f"{ASSETS}/logo_puc_campinas.png", Inches(0.55), Inches(0.4), height=Inches(0.72))
box, tf = add_textbox(s, Inches(1.35), Inches(0.42), Inches(4.5), Inches(0.75), anchor=MSO_ANCHOR.MIDDLE)
for i, line in enumerate(["PONTIFÍCIA", "UNIVERSIDADE", "CATÓLICA DE CAMPINAS"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.0
    _add_runs_markup(p, line, 10.5, BLACK, font=SANS)
    for r in p.runs:
        r.font.bold = True

box, tf = add_textbox(s, SLIDE_W - Inches(3.3), Inches(0.5), Inches(2.9), Inches(0.4))
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.RIGHT
_add_runs_markup(p, DATA_LABEL, 12, BLACK, font=SANS)

box, tf = add_textbox(s, Inches(0.7), Inches(2.3), Inches(11.9), Inches(2.9), anchor=MSO_ANCHOR.TOP)
p = tf.paragraphs[0]
p.line_spacing = 1.12
_add_runs_markup(
    p,
    "Um Aplicativo Móvel para Gerenciamento de Treinos com Detecção "
    "da Qualidade de Execução de Exercícios via Visão Computacional",
    34, NAVY, font=SERIF,
)
for r in p.runs:
    r.font.bold = False

box, tf = add_textbox(s, Inches(0.7), Inches(6.15), Inches(4.0), Inches(0.9))
p = tf.paragraphs[0]
_add_runs_markup(p, "Orientador", 13, BLACK, font=SANS)
for r in p.runs:
    r.font.bold = True
p2 = tf.add_paragraph()
_add_runs_markup(p2, "Antonio Paulo Salgado Forster", 14, BLUE, font=SANS)

box, tf = add_textbox(s, Inches(5.6), Inches(6.15), Inches(4.0), Inches(0.9))
p = tf.paragraphs[0]
_add_runs_markup(p, "Aluno", 13, BLACK, font=SANS)
for r in p.runs:
    r.font.bold = True
p2 = tf.add_paragraph()
_add_runs_markup(p2, "Fernando Bordin Correa", 14, BLUE, font=SANS)

box, tf = add_textbox(s, SLIDE_W - Inches(1.3), SLIDE_H - Inches(0.65), Inches(1.0), Inches(0.4))
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.RIGHT
_add_runs_markup(p, "TCC", 11, GRAY, font=SANS)


# ============================================================================
# SLIDE 2 — Sumário
# ============================================================================
s, bt = content_slide("Sumário", 2)
toc_left = [
    ("I", "Introdução", "3"),
    ("II", "Objetivos", "4"),
    ("III", "Revisão Tecnológica", "5"),
    ("IV", "Arquitetura do Sistema", "8"),
    ("V", "Detecção de Execução de Exercícios", "12"),
]
toc_right = [
    ("VI", "Implantação em Nuvem", "15"),
    ("VII", "Metodologia de Avaliação", "16"),
    ("VIII", "Resultados e Limitações", "17"),
    ("IX", "Trabalhos Futuros", "19"),
    ("X", "Conclusão", "20"),
]
for col, data in [(0, toc_left), (1, toc_right)]:
    left = MARGIN_L + Inches(col * 6.1)
    box, tf = add_textbox(s, left, Inches(2.0), Inches(5.6), Inches(4.6))
    for i, (num, label, pg) in enumerate(data):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(22)
        r1 = p.add_run()
        _set_run(r1, f"{num}   ", 15, BLUE, bold=True, font=SANS)
        r2 = p.add_run()
        _set_run(r2, label, 15, BLACK, font=SANS)
        r3 = p.add_run()
        _set_run(r3, f"\t{pg}", 15, GRAY, font=SANS)


# ============================================================================
# SLIDE 3 — Introdução: O Problema
# ============================================================================
s, bt = content_slide("Introdução", 3, subheading="O Problema")
add_bullets(s, [
    "Execução incorreta de exercícios resistidos é causa comum de **lesão** e de **baixa efetividade** do treino",
    "Praticantes iniciantes raramente têm um **profissional supervisionando cada repetição**",
    "A avaliação hoje depende de autoavaliação ou de vídeos genéricos, que **não analisam a execução específica do usuário**",
    "A popularização de câmeras de boa qualidade em _smartphones_ e o amadurecimento de bibliotecas de estimativa de pose em tempo real tornaram viável um sistema acessível, **sem sensores vestíveis**",
], top=bt, size=18, space_after=18)


# ============================================================================
# SLIDE 4 — Objetivos
# ============================================================================
s, bt = content_slide("Objetivos", 4, subheading="Objetivo Geral")
add_bullets(s, [
    "Desenvolver e avaliar um sistema mobile completo de apoio ao treino de musculação, com um módulo de detecção automática da qualidade de execução baseado em visão computacional",
], top=bt, size=17, height=Inches(1.1), space_after=0)

add_subheading(s, "Objetivos Específicos", top=Inches(3.15))
add_bullets(s, [
    "App móvel Android (React Native/Expo) para planejamento, execução e histórico de treinos",
    "Backend que centralize a persistência de dados e hospede o serviço de análise de vídeo",
    "Pipeline de extração de características a partir de pose humana + classificador correto/incorreto para mais de um exercício",
    "Implantação em um ambiente de nuvem, validada fora do ambiente de desenvolvimento local",
    "Identificar e documentar as limitações do modelo atual, propondo direções para superá-las",
], top=Inches(3.75), size=15.5, space_after=10)


# ============================================================================
# SLIDE 5 — Revisão Tecnológica: Sensores vs. Visão Computacional
# ============================================================================
s, bt = content_slide("Revisão Tecnológica", 5, subheading="Sensores Inerciais — Velloso et al. (2013)")
add_bullets(s, [
    "4 IMUs (acelerômetro + giroscópio + magnetômetro) fixadas em luva, bracelete, cinto lombar e halter",
    "Rosca bíceps unilateral, **5 classes**: 1 execução correta + 4 erros comuns",
    "Random Forest: **98,03% de acurácia** em validação cruzada 10-fold",
    "Limitação prática: exige **instrumentar o usuário** com múltiplos sensores — inviabiliza o uso espontâneo numa academia comum",
], top=bt, size=15.5, space_after=10, height=Inches(2.6))

add_subheading(s, "Visão Computacional", top=Inches(4.75))
add_bullets(s, [
    "Dispensa qualquer instrumentação além de uma câmera comum",
    "OpenPose (_bottom-up_, 18/25 pontos) vs. **MediaPipe/BlazePose** (33 pontos, até 75× mais rápido em CPU móvel) — usado neste trabalho",
    "Posição de câmera também afeta a estimativa de pose (Aguilar-Ortega et al., 2023): diferença **estatisticamente significativa** entre ângulos de câmera",
], top=Inches(5.3), size=15.5, space_after=8, height=Inches(2.0))


# ============================================================================
# SLIDE 6 — Revisão Tecnológica: Sistemas de Correção de Exercícios
# ============================================================================
s, bt = content_slide("Revisão Tecnológica", 6, subheading="Pose Trainer — Chen e Yang (2020)")
add_bullets(s, [
    "OpenPose + regras geométricas manuais + comparação por _Dynamic Time Warping_ (DTW) / kNN contra trajetórias de referência",
    "4 exercícios com halteres; F1-score entre **0,73 e 1,00**",
    "**Origem do dataset de treino usado neste trabalho** (elevação frontal e rosca bíceps)",
], top=bt, size=16, space_after=12, height=Inches(2.3))

add_subheading(s, "Trabalhos Recentes (2021+)", top=Inches(4.3))
add_bullets(s, [
    "Parmar, Gharat e Rhodin (2022): avaliação _self-supervised_ de erro de forma, direto do vídeo, sem regras manuais",
    "Dittakavi et al. (2022) — Pose Tutor: sistema **explicável**, aponta qual articulação causou o erro",
], top=Inches(4.85), size=16, space_after=10, height=Inches(1.7))


# ============================================================================
# SLIDE 7 — Posicionamento deste Trabalho
# ============================================================================
s, bt = content_slide("Posicionamento deste Trabalho", 7)
rows = [
    ("Eixo", "Velloso et al.", "Pose Trainer", "Este trabalho"),
    ("Sensor", "IMU vestível", "Câmera de smartphone", "Câmera de smartphone"),
    ("Estratégia de decisão", "Classificador sobre séries temporais", "Regras + DTW/kNN", "Classificador (RF/SVM) sobre o quadro de pico"),
    ("Feedback ao usuário", "Classe de erro (A–E)", "Qual erro geométrico", "Só correto/incorreto + confiança"),
    ("Origem do dataset", "Coletado pelos autores", "Coletado pelos autores", "Reaproveitado de terceiros (Pose Trainer)"),
    ("Restrição computacional", "Não se aplica", "Ambiente de pesquisa", "Nuvem sem GPU (Azure App Service B1)"),
]
n_rows, n_cols = len(rows), 4
tbl_w, tbl_h = Inches(11.9), Inches(4.7)
gfx = s.shapes.add_table(n_rows, n_cols, MARGIN_L, bt + Inches(0.1), tbl_w, tbl_h)
table = gfx.table
col_widths = [Inches(2.5), Inches(3.0), Inches(3.0), Inches(3.4)]
for ci, cw in enumerate(col_widths):
    table.columns[ci].width = cw
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        set_cell_text(table.cell(ri, ci), text, size=13.5 if ri == 0 else 13,
                      bold=(ri == 0), align=PP_ALIGN.CENTER if ri == 0 else PP_ALIGN.LEFT)
style_table(table, n_cols)
box, tf = add_textbox(s, MARGIN_L, bt + Inches(4.95), CONTENT_W, Inches(0.4))
p = tf.paragraphs[0]
_add_runs_markup(p, "_Tabela completa com 8 eixos de comparação no artigo (Seção 2.4)._", 12, GRAY, font=SANS)


# ============================================================================
# SLIDE 8 — Arquitetura do Sistema
# ============================================================================
s, bt = content_slide("Arquitetura do Sistema", 8)
add_bullets(s, [
    "3 componentes: **app móvel** (cliente), **backend** (HTTP + WebSocket) e **módulo de ML** — executado no mesmo processo do backend, não como serviço separado",
], top=bt, size=16, height=Inches(0.9), space_after=0)
add_picture_fit(s, f"{ASSETS}/diagrama_arquitetura.jpg", MARGIN_L, bt + Inches(0.85),
                 CONTENT_W, Inches(5.2))


# ============================================================================
# SLIDE 9 — Aplicativo Móvel
# ============================================================================
s, bt = content_slide("Aplicativo Móvel", 9, subheading="React Native + Expo (Android)")
add_bullets(s, [
    "Autenticação de usuário",
    "Cadastro de exercícios e montagem de fichas de treino (séries, repetições, carga)",
    "Execução de sessões com cronômetro",
    "Histórico e estatísticas mensais de treino",
    "Tela do detector: grava ou seleciona um vídeo para análise de execução",
], top=bt, size=18, space_after=16)
add_subheading(s, "Comunicação com o Backend", top=Inches(6.2), size=15)
add_bullets(s, [
    "HTTP (API REST, _token_ JWT) e, no detector, também WebSocket",
], top=Inches(6.65), size=14.5, height=Inches(0.5), space_after=0)


# ============================================================================
# SLIDE 10 — Diagrama de Casos de Uso
# ============================================================================
s = new_slide()
add_title(s, "Diagrama de Casos de Uso", top=Inches(0.35))
add_page_chrome(s, 10)
add_picture_fit(s, f"{ASSETS}/diagrama_use_cases.jpg", MARGIN_L, Inches(1.05), CONTENT_W, Inches(6.15))


# ============================================================================
# SLIDE 11 — Backend
# ============================================================================
s, bt = content_slide("Backend", 11)
add_bullets(s, [
    "**FastAPI** (Python) — rotas assíncronas (relevante para o processamento de vídeo) + documentação automática",
    "**SQLAlchemy** como ORM — SQLite em desenvolvimento, **PostgreSQL em produção**; migrações com Alembic",
    "Autenticação com **JWT** (_python-jose_) e senhas com _hash_ **bcrypt**",
    "Modelo de dados: usuários, exercícios, treinos e sessões — cada série com o registro real de peso/repetições executados",
], top=bt, size=18, space_after=18)


# ============================================================================
# SLIDE 12 — Detector: Estimativa de Pose
# ============================================================================
s, bt = content_slide("Detecção de Execução de Exercícios", 12, subheading="Estimativa de Pose e Extração de Características")
add_bullets(s, [
    "Cada quadro processado com **MediaPipe Pose** (via OpenCV) → 33 pontos-chave 2D",
    "Ângulos articulares + alturas relativas do punho, **normalizados** pela distância ombro–quadril do próprio usuário",
    "Cada exercício tem seu próprio conjunto de características e sua própria regra de \"repouso\" (delimita início/fim de cada repetição)",
], top=bt, size=16.5, space_after=12, height=Inches(2.6))

add_subheading(s, "Características por Exercício", top=Inches(4.75))
rows = [
    ("Exercício", "Características"),
    ("Elevação Frontal", "7 — ângulo de braço e cotovelo (cada lado), altura de punho, assimetria"),
    ("Rosca Bíceps", "3 — ângulo de cotovelo, ângulo de braço, altura de punho (lado mais visível)"),
]
gfx = s.shapes.add_table(3, 2, MARGIN_L, Inches(5.25), Inches(11.9), Inches(1.6))
table = gfx.table
table.columns[0].width = Inches(3.2)
table.columns[1].width = Inches(8.7)
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        set_cell_text(table.cell(ri, ci), text, size=13.5 if ri == 0 else 13, bold=(ri == 0))
style_table(table, 2)


# ============================================================================
# SLIDE 13 — Detector: Classificação Supervisionada
# ============================================================================
s, bt = content_slide("Detecção de Execução de Exercícios", 13, subheading="Classificação Supervisionada por Exercício")
add_bullets(s, [
    "Inferência: classifica o **quadro de pico** da repetição — o de maior ângulo de braço (elevação frontal) ou maior flexão de cotovelo (rosca bíceps)",
    "**Elevação Frontal:** Random Forest fixo",
    "**Rosca Bíceps:** escolha automática entre Random Forest, Gradient Boosting e SVM (validação cruzada 5-fold, F1 macro) → modelo atual é um **SVM com kernel RBF**",
], top=bt, size=17, space_after=16, height=Inches(3.2))

box, tf = add_textbox(s, MARGIN_L, Inches(5.3), CONTENT_W, Inches(1.5))
p = tf.paragraphs[0]
_add_runs_markup(p, "Quadro de pico, formalmente:", 14, BLACK, font=SANS)
p2 = tf.add_paragraph()
p2.space_before = Pt(8)
r = p2.add_run()
_set_run(r, "quadro classificado = arg maxₜ iₜ", 20, NAVY, bold=True, font=SERIF)
p3 = tf.add_paragraph()
p3.space_before = Pt(6)
_add_runs_markup(p3, "_iₜ = métrica de intensidade do exercício no quadro t da repetição_", 13, GRAY, font=SANS, base_italic=True)


# ============================================================================
# SLIDE 14 — Dataset de Treino
# ============================================================================
s, bt = content_slide("Dataset de Treino", 14)
add_bullets(s, [
    "Dataset público do projeto **Pose Trainer**, extraído via OpenPose (formato BODY_18)",
    "Cada quadro válido do vídeo vira uma linha independente — **não há segmentação por repetição** no treino",
], top=bt, size=16.5, space_after=12, height=Inches(1.6))

rows = [
    ("Exercício", "Vídeos (bom/ruim)", "Amostras (linhas)", "Correto/Incorreto"),
    ("Elevação Frontal", "15 / 13", "626", "343 / 283"),
    ("Rosca Bíceps", "6 / 10", "1158", "465 / 693"),
]
gfx = s.shapes.add_table(3, 4, MARGIN_L, Inches(3.65), Inches(11.9), Inches(1.6))
table = gfx.table
for ci, cw in enumerate([Inches(3.4), Inches(2.8), Inches(2.8), Inches(2.9)]):
    table.columns[ci].width = cw
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        set_cell_text(table.cell(ri, ci), text, size=15 if ri == 0 else 14.5,
                      bold=(ri == 0), align=PP_ALIGN.CENTER)
style_table(table, 4)

add_bullets(s, [
    "Vídeos com execuções rotuladas como \"correta\" ou \"incorreta\" (ex.: flexão do cotovelo na elevação frontal, uso de momento do corpo na rosca bíceps)",
], top=Inches(5.7), size=15, height=Inches(1.0), space_after=0)


# ============================================================================
# SLIDE 15 — Implantação em Nuvem
# ============================================================================
s, bt = content_slide("Implantação em Nuvem", 15)
add_bullets(s, [
    "Backend conteinerizado com **Docker** e implantado na **Microsoft Azure**:",
    (1, "Azure Container Registry — armazena a imagem"),
    (1, "Azure App Service, camada B1 — executa o container (sem GPU)"),
    (1, "Azure Database for PostgreSQL, camada B1ms — persistência"),
    "App móvel compilado como pacote Android autônomo (**EAS Build**, perfil _preview_), já apontando para o backend na nuvem",
], top=bt, size=17, space_after=14)


# ============================================================================
# SLIDE 16 — Metodologia de Avaliação
# ============================================================================
s, bt = content_slide("Metodologia de Avaliação", 16)
add_bullets(s, [
    "Particionamento treino/teste (80/20) e validação cruzada 5-fold feitos **por quadro**, não por vídeo",
    "O identificador do vídeo de origem **não é preservado** no CSV — impossível garantir que quadros de um mesmo vídeo fiquem só de um lado da partição",
    "**Métricas reportadas:** acurácia (teste + validação cruzada) para os dois exercícios; F1 macro (validação cruzada) adicional na rosca bíceps",
], top=bt, size=18, space_after=20)


# ============================================================================
# SLIDE 17 — Resultados
# ============================================================================
s, bt = content_slide("Resultados", 17)
add_bullets(s, [
    "Sistema completo implantado e funcional, testado a partir de um _build_ Android real, fora do ambiente de desenvolvimento",
], top=bt, size=17, height=Inches(0.9), space_after=0)

rows = [
    ("Exercício", "Acurácia (teste)", "Validação Cruzada", "Modelo"),
    ("Elevação Frontal", "79%", "83% ± 3% (acurácia)", "Random Forest"),
    ("Rosca Bíceps", "85%", "F1 macro 0,85", "SVM (kernel RBF)"),
]
gfx = s.shapes.add_table(3, 4, MARGIN_L, Inches(3.0), Inches(11.9), Inches(1.8))
table = gfx.table
for ci, cw in enumerate([Inches(3.0), Inches(2.9), Inches(3.3), Inches(2.7)]):
    table.columns[ci].width = cw
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        set_cell_text(table.cell(ri, ci), text, size=16 if ri == 0 else 16,
                      bold=(ri == 0), align=PP_ALIGN.CENTER)
style_table(table, 4, header_size=15, body_size=15.5)


# ============================================================================
# SLIDE 18 — Limitações
# ============================================================================
s, bt = content_slide("Limitações", 18)
add_bullets(s, [
    "Acurácias são um **limite superior otimista**: _split_ por quadro (não por vídeo) provavelmente vaza entre treino e teste",
    "**Mecanismo causal central:** a inferência sempre classifica o quadro de maior amplitude da repetição, mas o treino não seleciona picos — dataset cobre só até ~110°, execuções reais chegam a 172°",
    "Observação preliminar **não controlada**: câmera operada por terceiro erra mais que autorregistro (plausibilidade externa em Aguilar-Ortega et al., 2023)",
    "Descasamento de limiar de repouso na rosca bíceps (155° na curadoria do dataset vs. 130° em produção)",
    "Sem matriz de confusão nem métricas por classe; dataset da rosca bíceps desbalanceado (465/693)",
    "Segmentação de repetições (detecção de repouso) nunca avaliada isoladamente",
], top=bt, size=15.5, space_after=12)


# ============================================================================
# SLIDE 19 — Trabalhos Futuros
# ============================================================================
s, bt = content_slide("Trabalhos Futuros", 19)
add_bullets(s, [
    "**Metodologia de avaliação:** particionamento por vídeo (GroupKFold); alinhar amostragem de treino ao quadro de pico; unificar limiares de repouso; avaliar a segmentação de repetições; métricas por classe e matriz de confusão; experimento controlado de posicionamento de câmera",
    "**Dados e modelo:** ampliar o dataset com vídeos próprios (amplitude + posicionamento de câmera); extrair dataset direto do MediaPipe, unificando o estimador usado em treino e produção",
    "**Produto:** novos exercícios; validação de usabilidade com um grupo maior de usuários reais",
], top=bt, size=17, space_after=18)


# ============================================================================
# SLIDE 20 — Conclusão
# ============================================================================
s, bt = content_slide("Conclusão", 20)
add_bullets(s, [
    "App móvel completo de gerenciamento de treinos com um módulo de detecção da qualidade de execução baseado em estimativa de pose e aprendizado de máquina supervisionado",
    "Cobertura de ponta a ponta: da arquitetura do sistema até a implantação em um ambiente de nuvem real",
    "Principal lição do desenvolvimento: os maiores riscos técnicos não estão só na escolha do algoritmo de classificação, mas na **representatividade do dataset de treino** frente à variabilidade real de execução dos usuários",
], top=bt, size=18, space_after=20)


# ============================================================================
# SLIDE 21 — Agradecimentos
# ============================================================================
s = new_slide()
box, tf = add_textbox(s, Inches(1.2), Inches(2.6), Inches(10.9), Inches(3.5))
p = tf.paragraphs[0]
_add_runs_markup(p, "Agradecimentos", 38, NAVY, font=SERIF)
for r in p.runs:
    r.font.bold = True
items = [
    "Ao meu orientador, Prof. Antonio Paulo Salgado Forster, pela orientação ao longo do trabalho;",
    "Aos professores e colegas que acompanharam o desenvolvimento deste projeto;",
    "À minha família, pelo apoio constante.",
]
box2, tf2 = add_textbox(s, Inches(1.2), Inches(3.9), Inches(10.5), Inches(2.5))
for i, it in enumerate(items):
    p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
    p.space_after = Pt(14)
    _add_runs_markup(p, it, 18, BLACK, font=SANS)
add_page_chrome(s, 21)

box, tf = add_textbox(s, Inches(1.2), Inches(6.4), Inches(10.9), Inches(0.8))
p = tf.paragraphs[0]
_add_runs_markup(p, "Obrigado! Perguntas?", 22, BLUE, font=SANS)
for r in p.runs:
    r.font.bold = True


prs.save("Apresentacao_TCC.pptx")
print(f"Gerados {len(prs.slides.__iter__.__self__._sldIdLst)} slides -> Apresentacao_TCC.pptx")
